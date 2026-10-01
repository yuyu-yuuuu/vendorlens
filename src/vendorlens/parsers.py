from __future__ import annotations

import hashlib
import io
import re
from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader

from vendorlens.domain import ParsedChunk, SourceLocation


def _chunk(text: str, location: SourceLocation, **metadata) -> ParsedChunk:
    normalized = text.strip()
    return ParsedChunk(
        text=normalized,
        location=location,
        text_hash=hashlib.sha256(normalized.encode("utf-8")).hexdigest(),
        metadata=metadata,
    )


def parse_document(filename: str, content: bytes) -> list[ParsedChunk]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md"}:
        return _parse_text(content, markdown=suffix == ".md")
    if suffix == ".xlsx":
        return _parse_xlsx(content)
    if suffix == ".docx":
        return _parse_docx(content)
    if suffix == ".pdf":
        return _parse_pdf(content)
    raise ValueError(f"unsupported document type: {suffix or '(none)'}")


def _parse_text(content: bytes, markdown: bool) -> list[ParsedChunk]:
    text = content.decode("utf-8-sig")
    lines = text.splitlines()
    chunks: list[ParsedChunk] = []
    headings: list[str] = []
    buffer: list[str] = []
    start = 1

    def flush(end: int) -> None:
        nonlocal buffer, start
        value = "\n".join(buffer).strip()
        if value:
            chunks.append(
                _chunk(
                    value,
                    SourceLocation(
                        heading_path=tuple(headings), line_start=start, line_end=end
                    ),
                )
            )
        buffer = []

    for number, line in enumerate(lines, 1):
        heading = re.match(r"^(#{1,6})\s+(.+)$", line) if markdown else None
        if heading:
            flush(number - 1)
            level = len(heading.group(1))
            headings[:] = headings[: level - 1]
            headings.append(heading.group(2).strip())
            start = number + 1
        elif not line.strip():
            if buffer:
                flush(number - 1)
            start = number + 1
        else:
            if not buffer:
                start = number
            buffer.append(line)
    flush(len(lines))
    return chunks


def _parse_xlsx(content: bytes) -> list[ParsedChunk]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    chunks: list[ParsedChunk] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            values = ["" if cell.value is None else str(cell.value) for cell in row]
            if not any(value.strip() for value in values):
                continue
            start, end = row[0].coordinate, row[-1].coordinate
            chunks.append(
                _chunk(
                    " | ".join(values),
                    SourceLocation(sheet=sheet.title, cell_range=f"{start}:{end}"),
                )
            )
    return chunks


def _parse_docx(content: bytes) -> list[ParsedChunk]:
    document = DocxDocument(io.BytesIO(content))
    headings: list[str] = []
    chunks: list[ParsedChunk] = []
    for index, paragraph in enumerate(document.paragraphs, 1):
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style and paragraph.style.name.startswith("Heading"):
            level_text = paragraph.style.name.removeprefix("Heading").strip()
            level = int(level_text) if level_text.isdigit() else 1
            headings[:] = headings[: level - 1]
            headings.append(text)
            continue
        chunks.append(
            _chunk(text, SourceLocation(paragraph=index, heading_path=tuple(headings)))
        )
    for table_index, table in enumerate(document.tables, 1):
        for row_index, row in enumerate(table.rows, 1):
            text = " | ".join(cell.text.strip() for cell in row.cells)
            if text.strip(" |"):
                chunks.append(
                    _chunk(
                        text,
                        SourceLocation(heading_path=tuple(headings)),
                        table=table_index,
                        row=row_index,
                    )
                )
    return chunks


def _parse_pdf(content: bytes) -> list[ParsedChunk]:
    reader = PdfReader(io.BytesIO(content))
    chunks: list[ParsedChunk] = []
    for page_number, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            chunks.append(_chunk(text, SourceLocation(page=page_number)))
    if not chunks:
        raise ValueError("pdf contains no extractable text; OCR is required")
    return chunks
